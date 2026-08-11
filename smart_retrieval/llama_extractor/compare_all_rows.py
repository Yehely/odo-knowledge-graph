"""
compare_all_rows.py – side-by-side, human-readable dump of what's already
in the dataset for each PMID vs. what Llama extracted for it. Complements
Comparison.py's aggregate statistics with per-article detail for manual
review.

Usage:
    python3 compare_all_rows.py --input labeled_data_from_llama.json
"""
import argparse
import json
import os

import openpyxl

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DATASET_PATH = os.path.join(os.path.dirname(os.path.dirname(SCRIPT_DIR)), "Final ODO Dataset_v2026-06-10.xlsx")


def get_col_index(header, name):
    try:
        return header.index(name)
    except ValueError:
        return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True, help="Llama extraction JSON to inspect")
    args = ap.parse_args()

    print("Loading Llama extraction data...")
    with open(args.input, "r", encoding="utf-8") as f:
        llama_data = json.load(f)

    print(f"Opening dataset: {DATASET_PATH}...")
    wb = openpyxl.load_workbook(DATASET_PATH, read_only=True, data_only=True)
    sheet = wb["Full Dataset"] if "Full Dataset" in wb.sheetnames else wb.active

    header = [str(cell.value) for cell in sheet[1]]
    col_pmid = get_col_index(header, "pubmed_id")
    col_drug = get_col_index(header, "chembl_chemical_entity_key")
    col_target = get_col_index(header, "target_name")
    col_val = get_col_index(header, "endpoint_value")
    col_unit = get_col_index(header, "unit_of_measurement")
    col_type = get_col_index(header, "endpoint")

    unique_pmids = list({str(item.get("pmid")) for item in llama_data})

    for pid in unique_pmids:
        print("\n" + "#" * 100)
        print(f"Article PMID: {pid}")
        print("#" * 100)

        print("\nExisting rows in the dataset for this PMID:")
        excel_count = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[col_pmid] and str(row[col_pmid]).split(".")[0] == pid:
                excel_count += 1
                drug_val = row[col_drug] if row[col_drug] is not None else "N/A"
                print(f"   [DB] Drug: {str(drug_val):<15} | Target: {str(row[col_target]):<20} "
                      f"| Value: {str(row[col_val]):<6} {str(row[col_unit]):<4} | Type: {row[col_type]}")

        if excel_count == 0:
            print("   No existing rows found for this PMID in the dataset.")

        print("\nLlama extractions:")
        llama_findings = [i for i in llama_data if str(i.get("pmid")) == pid]

        for item in llama_findings:
            drug = item.get("drug_name", "N/A")
            target = item.get("target_name", "N/A")
            val = item.get("value_numeric", "N/A")
            unit = item.get("units", "N/A")
            m_type = item.get("measure_type", "N/A")

            print(f"   [Llama] Drug: {drug:<15} | Target: {target:<20} | Value: {val:<6} {unit:<4} | Type: {m_type}")
            source = item.get("source_sentence", "")
            if source:
                clean_source = source.replace("\n", " ").strip()
                print(f"           -> Source: \"{clean_source[:120]}...\"")

    wb.close()


if __name__ == "__main__":
    main()
